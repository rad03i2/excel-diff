from __future__ import annotations

import csv
import json
from dataclasses import asdict, dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

SUPPORTED = {".xlsx", ".xlsm", ".csv"}


class ExcelDiffError(ValueError):
    """Raised for invalid input or unsupported spreadsheet data."""


@dataclass(frozen=True)
class CellChange:
    sheet: str
    cell: str
    before: Any
    after: Any


@dataclass
class DiffReport:
    left: str
    right: str
    added_sheets: list[str]
    removed_sheets: list[str]
    changes: list[CellChange]

    @property
    def identical(self) -> bool:
        return not (self.added_sheets or self.removed_sheets or self.changes)

    def to_dict(self) -> dict[str, Any]:
        return {
            "left": self.left,
            "right": self.right,
            "identical": self.identical,
            "added_sheets": self.added_sheets,
            "removed_sheets": self.removed_sheets,
            "changes": [asdict(c) for c in self.changes],
        }

    def to_json(self) -> str:
        return json.dumps(self.to_dict(), ensure_ascii=False, indent=2, default=str)


def _safe_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def _load_csv(path: Path) -> dict[str, dict[str, Any]]:
    cells: dict[str, Any] = {}
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        for row_i, row in enumerate(csv.reader(handle), 1):
            for col_i, value in enumerate(row, 1):
                if value != "":
                    from openpyxl.utils import get_column_letter
                    cells[f"{get_column_letter(col_i)}{row_i}"] = value
    return {"CSV": cells}


def _load_excel(path: Path, *, formulas: bool) -> dict[str, dict[str, Any]]:
    try:
        wb = load_workbook(path, read_only=True, data_only=not formulas, keep_vba=path.suffix.lower() == ".xlsm")
    except Exception as exc:
        raise ExcelDiffError(f"Cannot read workbook: {path}: {exc}") from exc
    result: dict[str, dict[str, Any]] = {}
    try:
        for ws in wb.worksheets:
            cells: dict[str, Any] = {}
            for row in ws.iter_rows():
                for cell in row:
                    value = _safe_value(cell.value)
                    if value is not None:
                        cells[cell.coordinate] = value
            result[ws.title] = cells
    finally:
        wb.close()
    return result


def load_spreadsheet(path: str | Path, *, formulas: bool = False) -> dict[str, dict[str, Any]]:
    p = Path(path)
    if not p.is_file():
        raise ExcelDiffError(f"File not found: {p}")
    suffix = p.suffix.lower()
    if suffix not in SUPPORTED:
        raise ExcelDiffError(f"Unsupported format '{suffix}'. Supported: {', '.join(sorted(SUPPORTED))}")
    return _load_csv(p) if suffix == ".csv" else _load_excel(p, formulas=formulas)


def compare(left: str | Path, right: str | Path, *, formulas: bool = False, ignore_case: bool = False) -> DiffReport:
    left_data = load_spreadsheet(left, formulas=formulas)
    right_data = load_spreadsheet(right, formulas=formulas)
    left_names, right_names = set(left_data), set(right_data)
    changes: list[CellChange] = []

    def normalized(v: Any) -> Any:
        return v.casefold() if ignore_case and isinstance(v, str) else v

    for sheet in sorted(left_names & right_names):
        l_cells, r_cells = left_data[sheet], right_data[sheet]
        for coord in sorted(set(l_cells) | set(r_cells)):
            before, after = l_cells.get(coord), r_cells.get(coord)
            if normalized(before) != normalized(after):
                changes.append(CellChange(sheet, coord, before, after))

    return DiffReport(
        str(Path(left)), str(Path(right)), sorted(right_names - left_names),
        sorted(left_names - right_names), changes,
    )
